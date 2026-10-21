from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from services.export_service import ExportService
from services.matching_service import MatchRule, MatchingService, MultiMatchConfig
from utils.value_normalizers import NormalizationOptions, ValueKind


class MultiExportTests(unittest.TestCase):
    def test_exports_five_audit_sheets(self) -> None:
        left = pd.DataFrame({"Plaka": ["34ABC123", "06XYZ789"], "Tutar": [100, 200]})
        right = pd.DataFrame({"Araç": ["34 ABC 123", "35NO0001"], "Bedel": [100, 200]})
        config = MultiMatchConfig(
            (
                MatchRule("Plaka", "Araç", normalization=NormalizationOptions(kind=ValueKind.PLATE), primary=True),
                MatchRule("Tutar", "Bedel", normalization=NormalizationOptions(kind=ValueKind.DECIMAL)),
            )
        )
        result = MatchingService.compare_multi(left, right, config)
        with tempfile.TemporaryDirectory() as directory:
            path = ExportService.export_multi(result, Path(directory) / "sonuc.xlsx")
            workbook = load_workbook(path, read_only=True)
            self.assertEqual(
                workbook.sheetnames,
                [
                    "Eşleşen Veriler",
                    "Dosya 1 Eşleşmeyenler",
                    "Dosya 2 Eşleşmeyenler",
                    "Kısmi Eşleşenler",
                    "Eşleştirme Kuralları ve Özet",
                ],
            )
            matched = workbook["Eşleşen Veriler"]
            self.assertIn("Dosya 1 | Kaynak Satır", [cell.value for cell in matched[1]])
            self.assertIn("Dosya 2 | Kaynak Satır", [cell.value for cell in matched[1]])
            workbook.close()


if __name__ == "__main__":
    unittest.main()
