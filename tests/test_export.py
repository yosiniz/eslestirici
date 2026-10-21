from __future__ import annotations

import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from services.export_service import ExportService
from services.matching_service import MatchingService, OutputMode


class ExportServiceTests(unittest.TestCase):
    def setUp(self) -> None:
        reference = pd.DataFrame({"Plaka": ["34ABC123"]})
        parking = pd.DataFrame(
            {
                "Plaka": ["34 ABC 123", "06 XYZ 06"],
                "Giriş Tarihi": [datetime(2026, 1, 2, 10, 30), datetime(2026, 1, 3)],
                "Ücret": [125.5, 90],
            }
        )
        self.result = MatchingService.compare(reference, "Plaka", parking, "Plaka")

    def test_all_mode_adds_status_column(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "all.xlsx"
            ExportService.export(self.result, target, OutputMode.ALL, False)
            workbook = load_workbook(target, read_only=False, data_only=True)
            sheet = workbook["Sonuçlar"]
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual(headers[-1], "Eşleşme Durumu")
            self.assertEqual(sheet["D2"].value, "Eşleşti")
            self.assertEqual(sheet["D3"].value, "Eşleşmedi")
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertTrue(sheet.auto_filter.ref)
            workbook.close()

    def test_split_mode_creates_summary_and_two_data_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "split.xlsx"
            ExportService.export(self.result, target, OutputMode.MATCHED, True)
            workbook = load_workbook(target, read_only=False, data_only=True)
            self.assertEqual(workbook.sheetnames, ["Özet", "Eşleşenler", "Eşleşmeyenler"])
            self.assertEqual(workbook["Eşleşenler"].max_row, 2)
            self.assertEqual(workbook["Eşleşmeyenler"].max_row, 2)
            summary_values = {
                workbook["Özet"].cell(row=row, column=1).value:
                workbook["Özet"].cell(row=row, column=2).value
                for row in range(3, 9)
            }
            self.assertEqual(summary_values["Toplam kayıt sayısı"], 2)
            self.assertEqual(summary_values["Eşleşen kayıt sayısı"], 1)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
