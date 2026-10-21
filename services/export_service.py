"""Karşılaştırma sonuçlarını biçimlendirilmiş XLSX olarak dışa aktarır."""

from __future__ import annotations

import math
from collections.abc import Callable, Iterable
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services.matching_service import MatchResult, MultiMatchResult, OutputMode


class ExportServiceError(RuntimeError):
    """Kullanıcıya gösterilebilecek dışa aktarma hatası."""


class ExportService:
    """Bellek dostu write-only çalışma kitabı oluşturur."""

    HEADER_FILL = PatternFill("solid", fgColor="17365D")
    HEADER_FONT = Font(color="FFFFFF", bold=True)
    SUMMARY_LABEL_FILL = PatternFill("solid", fgColor="DCE6F1")
    SUMMARY_TITLE_FILL = PatternFill("solid", fgColor="17365D")
    MAX_EXCEL_ROWS = 1_048_576

    @classmethod
    def export_multi(
        cls,
        result: MultiMatchResult,
        output_path: str | Path,
        progress: Callable[[int, str], None] | None = None,
    ) -> Path:
        """Çoklu eşleştirme sonucunu beş denetim sayfasıyla dışa aktarır."""
        destination = Path(output_path)
        if destination.suffix.lower() != ".xlsx":
            destination = destination.with_suffix(".xlsx")
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook(write_only=True)
        try:
            sheets = [
                ("Eşleşen Veriler", result.matched_pairs),
                ("Dosya 1 Eşleşmeyenler", result.left_unmatched),
                ("Dosya 2 Eşleşmeyenler", result.right_unmatched),
                ("Kısmi Eşleşenler", result.partial_matches),
            ]
            for index, (title, dataframe) in enumerate(sheets):
                start = 5 + index * 20
                cls._append_dataframe_sheet(
                    workbook, title, dataframe, progress, start, start + 18
                )
            cls._append_multi_summary_sheet(workbook, result)
            if progress:
                progress(92, "Excel diske yazılıyor...")
            workbook.save(destination)
        except PermissionError as exc:
            raise ExportServiceError(
                "Çıktı dosyasına yazılamadı. Dosya Excel'de açık olabilir."
            ) from exc
        except ExportServiceError:
            raise
        except Exception as exc:
            raise ExportServiceError(f"Excel oluşturulurken hata oluştu: {exc}") from exc
        finally:
            workbook.close()
        if progress:
            progress(100, "Excel başarıyla oluşturuldu.")
        return destination

    @classmethod
    def _append_dataframe_sheet(
        cls,
        workbook: Workbook,
        title: str,
        dataframe: pd.DataFrame,
        progress: Callable[[int, str], None] | None,
        progress_start: int,
        progress_end: int,
    ) -> None:
        if len(dataframe) + 1 > cls.MAX_EXCEL_ROWS:
            raise ExportServiceError(f"{title} sayfası Excel satır sınırını aşıyor.")
        sheet = workbook.create_sheet(title=title)
        headers = [str(column) for column in dataframe.columns]
        sheet.freeze_panes = "A2"
        sheet.append([cls._header_cell(sheet, header) for header in headers])
        widths = [min(max(len(header) + 2, 10), 50) for header in headers]
        sample = dataframe.head(1_000)
        for row in sample.itertuples(index=False, name=None):
            for i, value in enumerate(row):
                widths[i] = min(max(widths[i], len(str(cls._clean_value(value) or "")) + 2), 50)
        for i, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(i)].width = width
        total = len(dataframe)
        for row_number, row in enumerate(dataframe.itertuples(index=False, name=None), 1):
            sheet.append([cls._value_cell(sheet, value) for value in row])
            if progress and (row_number % 5_000 == 0 or row_number == total):
                ratio = row_number / max(total, 1)
                progress(
                    progress_start + int((progress_end - progress_start) * ratio),
                    f"{title} yazılıyor... {row_number:,}/{total:,}",
                )
        if headers:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{total + 1}"

    @classmethod
    def _append_multi_summary_sheet(cls, workbook: Workbook, result: MultiMatchResult) -> None:
        sheet = workbook.create_sheet(title="Eşleştirme Kuralları ve Özet")
        sheet.freeze_panes = "A2"
        sheet.append([cls._header_cell(sheet, "Özet"), cls._header_cell(sheet, "Değer")])
        for label, value in result.summary.as_rows():
            sheet.append([cls._value_cell(sheet, label), cls._value_cell(sheet, value)])
        sheet.append([cls._value_cell(sheet, "Mantık"), cls._value_cell(sheet, result.config.logic.value.upper())])
        sheet.append([cls._value_cell(sheet, "Tekrar Modu"), cls._value_cell(sheet, result.config.duplicate_mode.value)])
        sheet.append([])
        rule_rows = result.rules_as_rows()
        headers = list(rule_rows[0]) if rule_rows else ["Kural"]
        sheet.append([cls._header_cell(sheet, header) for header in headers])
        for row in rule_rows:
            sheet.append([cls._value_cell(sheet, row[header]) for header in headers])
        for i, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(i)].width = min(max(len(header) + 4, 14), 35)

    @classmethod
    def export(
        cls,
        result: MatchResult,
        output_path: str | Path,
        mode: OutputMode,
        split_sheets: bool,
        progress: Callable[[int, str], None] | None = None,
    ) -> Path:
        destination = Path(output_path)
        if destination.suffix.lower() != ".xlsx":
            destination = destination.with_suffix(".xlsx")
        if destination.exists() and destination.is_dir():
            raise ExportServiceError("Çıktı yolu bir klasörü gösteriyor.")
        destination.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook(write_only=True)
        try:
            if progress:
                progress(5, "Excel oluşturuluyor...")

            if split_sheets:
                matched_indices = np.flatnonzero(result.match_mask)
                unmatched_indices = np.flatnonzero(~result.match_mask)
                cls._append_data_sheet(
                    workbook,
                    "Eşleşenler",
                    result.parking_data,
                    matched_indices,
                    match_mask=result.match_mask,
                    include_status=False,
                    progress=progress,
                    progress_start=10,
                    progress_end=48,
                )
                cls._append_data_sheet(
                    workbook,
                    "Eşleşmeyenler",
                    result.parking_data,
                    unmatched_indices,
                    match_mask=result.match_mask,
                    include_status=False,
                    progress=progress,
                    progress_start=48,
                    progress_end=88,
                )
                cls._append_summary_sheet(workbook, result)
            else:
                indices = result.indices_for_mode(mode)
                cls._append_data_sheet(
                    workbook,
                    "Sonuçlar",
                    result.parking_data,
                    indices,
                    match_mask=result.match_mask,
                    include_status=mode is OutputMode.ALL,
                    progress=progress,
                    progress_start=10,
                    progress_end=90,
                )

            if progress:
                progress(94, "Excel diske yazılıyor...")
            workbook.save(destination)
        except PermissionError as exc:
            raise ExportServiceError(
                "Çıktı dosyasına yazılamadı. Dosya Excel'de açık olabilir; kapatıp yeniden deneyin."
            ) from exc
        except OSError as exc:
            raise ExportServiceError(f"Çıktı dosyası yazılamadı: {exc}") from exc
        except ExportServiceError:
            raise
        except Exception as exc:
            raise ExportServiceError(f"Excel oluşturulurken hata oluştu: {exc}") from exc
        finally:
            workbook.close()

        if progress:
            progress(100, "Excel başarıyla oluşturuldu.")
        return destination

    @classmethod
    def _append_data_sheet(
        cls,
        workbook: Workbook,
        title: str,
        dataframe: pd.DataFrame,
        indices: np.ndarray,
        match_mask: np.ndarray,
        include_status: bool,
        progress: Callable[[int, str], None] | None,
        progress_start: int,
        progress_end: int,
    ) -> None:
        row_count = int(len(indices))
        if row_count + 1 > cls.MAX_EXCEL_ROWS:
            raise ExportServiceError(
                f"{title} sayfası Excel'in satır sınırını aşıyor ({row_count:,} kayıt)."
            )

        sheet = workbook.create_sheet(title=title)
        headers = [str(column) for column in dataframe.columns]
        status_header = cls._unique_status_header(headers)
        output_headers = headers + ([status_header] if include_status else [])
        sheet.freeze_panes = "A2"
        sheet.append([cls._header_cell(sheet, header) for header in output_headers])

        widths = cls._estimate_widths(dataframe, indices, output_headers, include_status)
        for column_number, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column_number)].width = width

        report_every = max(2_000, row_count // 100) if row_count else 2_000
        written = 0
        chunk_size = 10_000
        for chunk_start in range(0, row_count, chunk_size):
            chunk_indices = indices[chunk_start : chunk_start + chunk_size]
            chunk = dataframe.iloc[chunk_indices]
            rows = chunk.itertuples(index=False, name=None)
            for position, source_values in zip(chunk_indices, rows, strict=True):
                cells = [cls._value_cell(sheet, value) for value in source_values]
                if include_status:
                    cells.append(
                        cls._value_cell(
                            sheet,
                            "Eşleşti" if bool(match_mask[int(position)]) else "Eşleşmedi",
                        )
                    )
                sheet.append(cells)
                written += 1

            if progress and (written % report_every == 0 or written == row_count):
                ratio = written / max(row_count, 1)
                value = progress_start + int((progress_end - progress_start) * ratio)
                progress(value, f"{title} sayfası yazılıyor... {written:,}/{row_count:,}")

        last_column = get_column_letter(max(len(output_headers), 1))
        sheet.auto_filter.ref = f"A1:{last_column}{row_count + 1}"

    @classmethod
    def _append_summary_sheet(cls, workbook: Workbook, result: MatchResult) -> None:
        sheet = workbook.create_sheet(title="Özet", index=0)
        sheet.freeze_panes = "A3"
        title_cell = WriteOnlyCell(sheet, value="Otopark Plaka Karşılaştırma Özeti")
        title_cell.font = Font(color="FFFFFF", bold=True, size=14)
        title_cell.fill = cls.SUMMARY_TITLE_FILL
        title_cell.alignment = Alignment(horizontal="left")
        empty_title = WriteOnlyCell(sheet, value=None)
        empty_title.fill = cls.SUMMARY_TITLE_FILL
        sheet.append([title_cell, empty_title])
        sheet.append([])

        for label, value in result.summary.as_rows():
            label_cell = WriteOnlyCell(sheet, value=label)
            label_cell.font = Font(bold=True, color="17365D")
            label_cell.fill = cls.SUMMARY_LABEL_FILL
            value_cell = WriteOnlyCell(sheet, value=value)
            value_cell.number_format = "#,##0"
            sheet.append([label_cell, value_cell])

        sheet.column_dimensions["A"].width = 42
        sheet.column_dimensions["B"].width = 18

    @classmethod
    def _header_cell(cls, sheet: Any, value: str) -> WriteOnlyCell:
        cell = WriteOnlyCell(sheet, value=value)
        cell.font = cls.HEADER_FONT
        cell.fill = cls.HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        return cell

    @classmethod
    def _value_cell(cls, sheet: Any, value: Any) -> WriteOnlyCell:
        clean_value = cls._clean_value(value)
        cell = WriteOnlyCell(sheet, value=clean_value)
        if isinstance(clean_value, datetime):
            cell.number_format = "dd.mm.yyyy hh:mm:ss"
        elif isinstance(clean_value, date):
            cell.number_format = "dd.mm.yyyy"
        elif isinstance(clean_value, time):
            cell.number_format = "hh:mm:ss"
        elif isinstance(clean_value, float):
            cell.number_format = "#,##0.00########"
        elif isinstance(clean_value, str):
            cell.data_type = "s"
        return cell

    @staticmethod
    def _clean_value(value: Any) -> Any:
        if value is None or value is pd.NA or value is pd.NaT:
            return None
        is_missing = False
        try:
            if not isinstance(value, (str, bytes)):
                is_missing = bool(pd.isna(value))
        except (TypeError, ValueError):
            is_missing = False
        if is_missing:
            return None

        if isinstance(value, np.generic):
            value = value.item()
        if isinstance(value, pd.Timestamp):
            value = value.to_pydatetime()
        if isinstance(value, datetime) and value.tzinfo is not None:
            value = value.replace(tzinfo=None)
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return None
        if isinstance(value, str):
            return ILLEGAL_CHARACTERS_RE.sub("", value)
        if isinstance(value, (list, tuple, dict, set)):
            return str(value)
        return value

    @classmethod
    def _estimate_widths(
        cls,
        dataframe: pd.DataFrame,
        indices: np.ndarray,
        output_headers: list[str],
        include_status: bool,
    ) -> list[float]:
        widths = [min(max(len(header) + 2, 10), 50) for header in output_headers]
        if len(indices) == 0:
            return widths

        sample_positions = indices[: min(len(indices), 1_000)]
        for position in sample_positions:
            values: Iterable[Any] = dataframe.iloc[int(position)].tolist()
            for column_index, value in enumerate(values):
                clean_value = cls._clean_value(value)
                if clean_value is None:
                    continue
                display_length = len(str(clean_value)) + 2
                widths[column_index] = min(max(widths[column_index], display_length), 50)
        if include_status:
            widths[-1] = max(widths[-1], 14)
        return widths

    @staticmethod
    def _unique_status_header(headers: list[str]) -> str:
        base = "Eşleşme Durumu"
        if base not in headers:
            return base
        suffix = 2
        while f"{base} ({suffix})" in headers:
            suffix += 1
        return f"{base} ({suffix})"
